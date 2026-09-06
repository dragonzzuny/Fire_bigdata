#!/usr/bin/env python
"""대본을 발표자료의 '슬라이드 노트' 로 옮긴다.

발표장에서 발표자 화면에 뜨는 것은 노트다. 그런데 노트가 비어 있었다 —
대본은 따로 종이로 들고 있어야 했다. 손에 든 종이와 화면이 따로 놀면
어느 쪽이 최신인지 알 수 없다.

노트를 손으로 옮겨 적지 않는다. docs/DEMO_SCRIPT.md 의 §2 를 읽어 장표마다
붙인다. 대본을 고치면 이 스크립트를 다시 돌리는 것으로 끝난다.

읽는 말과 하는 일을 노트에서도 나눠 적는다. 발표자 화면에서 그 둘이
섞여 있으면 긴장했을 때 지시문을 소리 내어 읽는다.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from pptx import Presentation                            # noqa: E402

from firebird.config import load_config                  # noqa: E402

SCRIPT = ROOT / "docs" / "DEMO_SCRIPT.md"

#: 블록 머리. 예: **⑦ 검증 · 2분 18초 · 장표 13·14·15**
HEAD = re.compile(r"^\*\*([①-⑪])\s*([^·*]+?)\s*·\s*([^·*]+?)\s*·\s*([^*]+?)\*\*\s*$")
VIDEO = re.compile(r"^\*\*⑥\s*시연\s*·\s*([^*]+?)\*\*\s*$")


def blocks() -> list[dict]:
    body = SCRIPT.read_text(encoding="utf-8")
    body = body.split("## 2. 실제로 말할 대본", 1)[1].split("## 3. 상세 대본", 1)[0]
    lines = body.splitlines()
    heads = [i for i, ln in enumerate(lines) if HEAD.match(ln) or VIDEO.match(ln)]
    out = []
    for j, i in enumerate(heads):
        end = heads[j + 1] if j + 1 < len(heads) else len(lines)
        tail = lines[i].split("장표", 1)
        slides = [int(x) for x in re.findall(r"\d+", tail[1])] if len(tail) > 1 else []
        out.append({"head": lines[i].strip("* "), "body": lines[i:end],
                    "slides": slides})
    return out


TURN = re.compile(r"^▶\s*하는 일\s*—\s*장표\s*\*{0,2}(\d+)")


def split_by_slide(block: dict) -> list[tuple[int, list[str]]]:
    """블록을 '장표 N 로 넘깁니다' 자리에서 쪼갠다.

    한 블록이 장표 셋을 걸치면 셋에 같은 전문을 붙이게 된다. 발표자
    화면에 지금 장표와 상관없는 말이 절반씩 섞이면 오히려 못 읽는다.
    """
    slides = block["slides"]
    parts, cur, who = [], [], slides[0] if slides else 0
    for ln in block["body"]:
        m = TURN.match(ln.strip())
        if m and cur:
            parts.append((who, cur))
            who, cur = int(m.group(1)), []
        cur.append(ln)
    if cur:
        parts.append((who, cur))
    # 대본이 말하지 않은 장표는 앞 조각에 딸려 둔다.
    seen = {w for w, _ in parts}
    for sl in slides:
        if sl not in seen and parts:
            parts.append((sl, parts[0][1]))
    return parts


def render(head: str, body: list[str]) -> str:
    """노트 한 장 분량. 읽는 말과 하는 일을 나눠 적는다."""
    say, doing, mode = [], [], None
    for ln in body:
        t = ln.strip()
        if t.startswith("▷"):
            mode = "say"
            continue
        if t.startswith("▶"):
            mode = "do"
            doing.append(re.sub(r"^▶\s*하는 일\s*—?\s*", "", t))
            continue
        if t.startswith(">"):
            v = t[1:].strip()
            if v:
                (say if mode != "do" else doing).append(v)
            continue
        if mode == "do" and t and not t.startswith(("|", "**", "`", "-")):
            doing.append(t)

    def clean(xs):
        return [re.sub(r"\*\*", "", x) for x in xs if x.strip()]

    parts = [head]
    if doing:
        parts += ["", "[하는 일]"] + [f"· {x}" for x in clean(doing)]
    if say:
        parts += ["", "[읽는다]"] + clean(say)
    return "\n".join(parts)


def main() -> int:
    cfg = load_config()
    decks = sorted(cfg.paths.outputs.glob("불씨예보_발표자료_*.pptx"))
    if not decks:
        print("발표자료가 없다. scripts/07_deck.py 를 먼저 돌려라.")
        return 1
    prs = Presentation(str(decks[0]))
    n = len(prs.slides._sldIdLst)

    notes: dict[int, list[str]] = {}
    for b in blocks():
        for sl, part in split_by_slide(b):
            if 1 <= sl <= n:
                notes.setdefault(sl, []).append(render(b["head"], part))

    # 영상이 대신하는 화면은 대본이 가리키지 않는다. 왜 안 넘기는지 적어 둔다.
    for sl in range(1, n + 1):
        if sl not in notes:
            notes[sl] = ["(영상이 대신하는 화면 — 영상이 끝난 뒤 3초씩 스친다)"]

    for i, slide in enumerate(prs.slides, 1):
        slide.notes_slide.notes_text_frame.text = "\n\n".join(notes[i])
    prs.save(decks[0])

    filled = sum(1 for i in range(1, n + 1)
                 if not notes[i][0].startswith("(영상"))
    print(f"슬라이드 노트 {n}장에 기록  (대본에서 온 것 {filled}장)")
    for i in range(1, n + 1):
        head = notes[i][0].splitlines()[0]
        print(f"  {i:2}  {head[:58]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
