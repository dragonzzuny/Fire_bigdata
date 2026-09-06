"""사실성 장치가 실제로 작동하는지 확인한다.

숫자를 지어내거나 조문을 바꾸는 것은 이 서비스에서 가장 큰 사고다.
막는 장치가 있다는 것만으로는 부족하고, 그 장치가 실제로 잡아내는지를
위조 입력으로 확인해야 한다.
"""
from __future__ import annotations

import sys
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from firebird import assistant as A, plans as PLN     # noqa: E402


DOC = """# 2026년 3월 10일 예방순찰 계획(안)

울산소방서

1. 관련
   가. 「소방기본법」 제10조
   나. 「화재의 예방 및 안전관리에 관한 법률」 제7조

2. 위 호와 관련하여 아래와 같이 예방순찰을 실시하고자 합니다.

| 출동관서 | 구역 | 거리 |
|---|---|---|
| 삼산119안전센터 | 2331_3456 | 4.4 km |
| 성남119안전센터 | 2329_3460 | 6.0 km |

전년 대비 포착률 68.5% (95% 신뢰구간 63.5~72.9%)

붙임 1. 순찰 대상 목록 1부.  끝.
"""


class TestPlanFidelity(unittest.TestCase):
    """다듬은 계획서가 원문의 수·조문을 그대로 담는지."""

    def test_문체만_바꾼_문서는_통과한다(self):
        cand = DOC.replace("실시하고자 합니다", "시행하고자 합니다")
        self.assertTrue(PLN.check_fidelity(DOC, cand)["ok"])

    def test_숫자를_반올림하면_막는다(self):
        r = PLN.check_fidelity(DOC, DOC.replace("68.5%", "약 70%"))
        self.assertFalse(r["ok"])
        self.assertIn("70", r["added_numbers"])

    def test_조문_번호를_바꾸면_막는다(self):
        r = PLN.check_fidelity(DOC, DOC.replace("제7조", "제17조"))
        self.assertFalse(r["ok"])
        self.assertIn("제17조", r["added_legal"])
        self.assertIn("제7조", r["dropped_legal"])

    def test_법령명을_바꾸면_막는다(self):
        r = PLN.check_fidelity(DOC, DOC.replace("소방기본법", "소방시설법"))
        self.assertFalse(r["ok"])

    def test_없던_별표를_붙이면_막는다(self):
        r = PLN.check_fidelity(DOC, DOC + "\n※ 별표 4 참조\n")
        self.assertFalse(r["ok"])

    def test_격자_번호를_바꾸면_막는다(self):
        r = PLN.check_fidelity(DOC, DOC.replace("2331_3456", "2331_3999"))
        self.assertFalse(r["ok"])

    def test_표_행이_통째로_빠지면_막는다(self):
        cand = DOC.replace("| 성남119안전센터 | 2329_3460 | 6.0 km |\n", "")
        r = PLN.check_fidelity(DOC, cand)
        self.assertFalse(r["ok"])
        self.assertTrue(r["dropped_numbers"])

    def test_숫자_하나만_빠져도_막는다(self):
        """표의 한 칸이 조용히 사라진 계획서는 틀린 계획서다."""
        r = PLN.check_fidelity(DOC, DOC.replace("4.4 km", ""))
        self.assertFalse(r["ok"])
        self.assertIn("4.4", r["dropped_numbers"])

    def test_음수_부호가_사라지면_막는다(self):
        a = "전년 대비 -3.5%p 감소"
        self.assertFalse(PLN.check_fidelity(a, "전년 대비 3.5%p 감소")["ok"])

    def test_낫표_없는_법령명_변조도_막는다(self):
        a = "근거: 소방기본법 시행규칙 제6조"
        self.assertFalse(PLN.check_fidelity(a, "근거: 소방시설법 시행규칙 제6조")["ok"])

    def test_한을_1개로_고치는_문체_교정은_통과한다(self):
        """공문투로 고치면 '한 격자'가 '1개 격자'가 된다. 사실은 바뀌지 않았다.

        이걸 위반으로 세면 다듬기가 한 번도 채택되지 않아 기능이 죽는다.
        실제로 그렇게 되어 있었고, 실측으로 확인해 규칙을 고쳤다.
        """
        a = "한 격자는 도보 5~7분 거리의 한 블록 범위입니다."
        b = "1개 격자는 도보 5~7분 거리의 1개 블록 범위임."
        r = PLN.check_fidelity(a, b)
        self.assertTrue(r["ok"], r)
        self.assertEqual(r["added_numbers_weighty"], [])

    def test_반올림은_막는다(self):
        a = "총 이동 41.4 km"
        r = PLN.check_fidelity(a, "총 이동 41 km")
        self.assertFalse(r["ok"])
        self.assertIn("41.4", r["dropped_numbers"])

    def test_자릿점_표기_차이는_넘긴다(self):
        a = "총 1,234 m 이동"
        b = "총 1234m 이동"
        self.assertTrue(PLN.check_fidelity(a, b)["ok"])


class TestAnswerGrounding(unittest.TestCase):
    """업무 도우미 답변이 검색된 자료 밖의 조문을 인용하지 않는지."""

    def _hits(self):
        return [(A.Doc(
            doc_id="d1", source="법령",
            title="화재의 예방 및 안전관리에 관한 법률 제7조(화재안전조사)",
            ref="화재의 예방 및 안전관리에 관한 법률 제7조",
            text="소방관서장은 화재안전조사를 실시할 수 있다. 별표 2 를 따른다. "
                 "위반한 자는 같은 법 제50조제1항에 따라 처벌한다."), 9.0)]

    def test_자료_안의_조문은_통과한다(self):
        ans = ("「화재의 예방 및 안전관리에 관한 법률」 제7조에 따라 조사할 수 "
               "있으며, 별표 2 를 따릅니다.")
        self.assertTrue(A.check_grounding(ans, self._hits())["ok"])

    def test_괘선으로_끊긴_원문도_찾아낸다(self):
        """별표 본문은 표 괘선이 낱말 중간에 박혀 온다. 그래도 대조되어야 한다."""
        g = A.check_grounding("같은 법 제50조제1항에 따릅니다.", self._hits())
        self.assertTrue(g["ok"], g["unsupported"])

    def test_없는_조문을_인용하면_잡아낸다(self):
        g = A.check_grounding("제31조의2 에 따릅니다.", self._hits())
        self.assertFalse(g["ok"])
        self.assertIn("조:제31조의2", g["unsupported"])

    def test_없는_법령을_인용하면_잡아낸다(self):
        g = A.check_grounding("「소방시설 설치 및 관리에 관한 법률」 제7조입니다.",
                              self._hits())
        self.assertFalse(g["ok"])

    def test_없는_별표를_인용하면_잡아낸다(self):
        g = A.check_grounding("별표 9 를 보십시오.", self._hits())
        self.assertFalse(g["ok"])

    def test_조문_본문이_없으면_항까지는_확인하지_못했다고_말한다(self):
        """자료에 그 조문의 본문이 없고 다른 문서의 인용만 있으면,
        그 항이 실재하는지 알 수 없다. 모르는 것을 위반으로 세면
        '확인했다'는 말의 값이 떨어진다. 모른다고 말하고 넘긴다."""
        g = A.check_grounding("같은 법 제50조제9항에 따릅니다.", self._hits())
        self.assertTrue(g["ok"])
        self.assertIn("조:제50조제9항", g["unverified"])

    def test_원문에_있는_항은_그대로_통과한다(self):
        g = A.check_grounding("같은 법 제50조제1항에 따릅니다.", self._hits())
        self.assertTrue(g["ok"])
        self.assertEqual(g["unverified"], [])

    def test_법령과_조문을_짝지어_본다(self):
        """자료에 「법A」와 (다른 법의) 제7조가 각각 있다고 해서
        「법A」 제7조가 근거가 되지는 않는다."""
        hits = [
            (A.Doc(doc_id="1", source="법령", title="소방기본법 제10조",
                   ref="소방기본법 제10조", text="소방용수시설의 설치 기준"), 9.0),
            (A.Doc(doc_id="2", source="법령",
                   title="화재의 예방 및 안전관리에 관한 법률 제7조",
                   ref="화재의 예방 및 안전관리에 관한 법률 제7조",
                   text="소방관서장은 화재안전조사를 실시할 수 있다"), 8.0)]
        self.assertTrue(A.check_grounding("「소방기본법」 제10조에 따릅니다.", hits)["ok"])
        bad = A.check_grounding("「소방기본법」 제7조에 따릅니다.", hits)
        self.assertFalse(bad["ok"])
        self.assertTrue(bad["unpaired"])

    def test_자료에_흩어진_조각으로_짝을_만들지_않는다(self):
        """어느 문서 본문이 다른 법을 언급하고, 그 문서가 제7조라고 해서
        '그 다른 법 제7조'가 근거가 되지는 않는다. 짝은 문서의 정체(제목·근거)
        또는 본문의 붙여 쓴 인용으로만 인정한다."""
        hits = [(A.Doc(
            doc_id="1", source="법령",
            title="화재의 예방 및 안전관리에 관한 법률 제7조(화재안전조사)",
            ref="화재의 예방 및 안전관리에 관한 법률 제7조",
            text="소방관서장은 「소방기본법」에 따른 소방활동과 별도로 "
                 "화재안전조사를 실시할 수 있다."), 9.0)]
        self.assertTrue(A.check_grounding(
            "「화재의 예방 및 안전관리에 관한 법률」 제7조입니다.", hits)["ok"])
        bad = A.check_grounding("「소방기본법」 제7조입니다.", hits)
        self.assertFalse(bad["ok"])
        self.assertTrue(bad["unpaired"])

    def test_공식_약칭은_위반이_아니다(self):
        """모델은 '화재예방법' 이라 쓰고 원문은 정식 명칭만 담고 있다.
        약칭을 위반으로 세면 경고가 쏟아져 진짜 위반이 묻힌다."""
        hits = [(A.Doc(
            doc_id="1", source="법령",
            title="화재의 예방 및 안전관리에 관한 법률 제7조",
            ref="화재의 예방 및 안전관리에 관한 법률 제7조",
            text="화재안전조사를 실시할 수 있다."), 9.0)]
        self.assertTrue(A.check_grounding("「화재예방법」 제7조입니다.", hits)["ok"])

    def test_항은_조문_본문의_동그라미_숫자로_대조한다(self):
        """조문 본문은 항을 ①②③ 으로 적는다. 없는 항은 걸러야 하고,
        있는 항을 위반으로 세면 안 된다."""
        hits = [(A.Doc(
            doc_id="1", source="법령", title="화재의 예방 및 안전관리에 관한 법률 제7조",
            ref="화재의 예방 및 안전관리에 관한 법률 제7조",
            text="① 소방관서장은 화재안전조사를 할 수 있다. "
                 "② 조사 항목은 대통령령으로 정한다."), 9.0)]
        self.assertTrue(A.check_grounding("제7조제2항에 따릅니다.", hits)["ok"])
        bad = A.check_grounding("제7조제9항에 따릅니다.", hits)
        self.assertFalse(bad["ok"])
        self.assertIn("조:제7조제9항", bad["unsupported"])

    def test_낫표_없이_적어도_짝을_본다(self):
        hits = [
            (A.Doc(doc_id="1", source="법령", title="소방기본법 제10조",
                   ref="소방기본법 제10조", text="소방용수시설"), 9.0),
            (A.Doc(doc_id="2", source="법령",
                   title="화재의 예방 및 안전관리에 관한 법률 제7조",
                   ref="화재의 예방 및 안전관리에 관한 법률 제7조",
                   text="화재안전조사"), 8.0)]
        self.assertFalse(A.check_grounding("소방기본법 제7조에 따릅니다.", hits)["ok"])

    def test_모른다고_답하면_통과한다(self):
        g = A.check_grounding("제공된 자료에서 확인되지 않습니다.", self._hits())
        self.assertTrue(g["ok"])


class TestLawTextCleaning(unittest.TestCase):
    """법제처 별표 본문의 표 괘선 제거."""

    def test_낱말_중간_괘선은_지운다(self):
        from firebird.lawdata import _strip_rules
        self.assertIn("제50조제1항", _strip_rules("같은 법 제50│조제1항에 따라"))

    def test_칸_경계_괘선은_공백으로(self):
        from firebird.lawdata import _strip_rules
        out = _strip_rules("일자 │ 조치사항 │ 보완일자")
        self.assertNotIn("│", out)
        self.assertIn("일자", out)
        self.assertIn("조치사항", out)


if __name__ == "__main__":
    unittest.main(verbosity=2)


class TestStationMatching(unittest.TestCase):
    """관서명 검색 결과가 정말 그 관서인지 가려내는 규칙.

    카카오 키워드 검색은 첫 결과를 그대로 믿으면 안 된다. '울주소방서' 로
    검색하면 '울산남울주소방서' 가 1순위로 오는데 이 둘은 다른 관서이며,
    실제로 그 오류 때문에 울주소방서 순찰 동선이 온산에서 출발했다.
    """

    def test_행정구역_접두사만_다르면_같은_관서(self):
        from firebird.stations import _score
        self.assertEqual(_score("남부소방서", "울산남부소방서"), 3)
        self.assertEqual(_score("삼산119안전센터", "울산 삼산119안전센터"), 3)

    def test_같은_부지_안의_시설도_받아들인다(self):
        from firebird.stations import _score
        self.assertEqual(_score("조치원소방서", "조치원소방서 전기차충전소"), 3)

    def test_앞에_글자가_붙으면_다른_관서다(self):
        from firebird.stations import _score
        self.assertLess(_score("울주소방서", "울산남울주소방서"), 3)
        self.assertLess(_score("세종소방서", "세종남부소방서"), 3)

    def test_전혀_다른_이름은_0점(self):
        from firebird.stations import _score
        self.assertEqual(_score("온산소방서", "울산남울주소방서"), 0)

    def test_관할에서_멀면_관할_중심으로_물러난다(self):
        """이름이 맞아도 좌표가 자기 관할 밖이면 다른 곳을 찾은 것이다."""
        import numpy as np
        import pandas as pd
        from firebird import stations as ST

        rng = np.random.default_rng(0)
        n = 60
        panel = pd.DataFrame({
            "grid_id": [f"g{i}" for i in range(n)],
            "center": ["가안전센터"] * (n // 2) + ["나안전센터"] * (n - n // 2),
            "lon": np.r_[129.30 + rng.normal(0, .01, n // 2),
                         129.40 + rng.normal(0, .01, n - n // 2)],
            "lat": np.r_[35.54 + rng.normal(0, .01, n // 2),
                         35.60 + rng.normal(0, .01, n - n // 2)],
            "pred": rng.gamma(1, 1, n),
            "target_total": rng.integers(0, 20, n),
            "fires": rng.poisson(.4, n),
        })

        class _Cfg:
            class paths:
                cache = Path("/nonexistent")

        # 관할에서 400km 떨어진 좌표를 받은 것으로 꾸민다
        real = ST.locate_stations
        ST.locate_stations = lambda names, cfg, **kw: pd.DataFrame({
            "name": ["가안전센터", "나안전센터"],
            "lon": [129.30, 126.98], "lat": [35.54, 37.57],
            "matched": [True, True], "address": ["울산 남구", "서울 중구"]})
        try:
            t = ST.station_table(panel, _Cfg(), level="center",
                                 city_label="울산광역시", allow_network=False)
        finally:
            ST.locate_stations = real

        row = t[t["name"] == "나안전센터"].iloc[0]
        self.assertFalse(bool(row["matched"]))
        self.assertIn("좌표 이상", str(row["coord_source"]))
        self.assertLess(abs(float(row["lat"]) - 35.60), 0.05,
                        "관할 중심으로 돌아와야 한다")
        ok = t[t["name"] == "가안전센터"].iloc[0]
        self.assertTrue(bool(ok["matched"]))


class TestBuildingLedgerApi(unittest.TestCase):
    """건축물대장 연계.

    키가 없는 환경에서도 아무것도 깨지지 않아야 하고, 응답이 오면 정확히
    읽어야 한다. 반쯤 채워진 값이 결재 문서에 들어가는 것이 최악이다.
    """

    OK_XML = """<response><header><resultCode>00</resultCode></header><body><items>
      <item><platPlc>울산 남구 달동 100</platPlc><newPlatPlc>울산 남구 삼산로 100</newPlatPlc>
      <bldNm>가나빌딩</bldNm><totArea>5230.5</totArea><archArea>820.2</archArea>
      <useAprDay>19980312</useAprDay><grndFlrCnt>8</grndFlrCnt>
      <mainPurpsCdNm>업무시설</mainPurpsCdNm></item>
      <item><platPlc>울산 남구 달동 101</platPlc><newPlatPlc>울산 남구 삼산로 102</newPlatPlc>
      <bldNm>다라상가</bldNm><totArea>1120.0</totArea><archArea>310.0</archArea>
      <useAprDay>20140820</useAprDay><grndFlrCnt>4</grndFlrCnt>
      <mainPurpsCdNm>근린생활시설</mainPurpsCdNm></item>
      </items><totalCount>2</totalCount></body></response>"""

    def test_표제부를_읽는다(self):
        from firebird import buildings as B
        rows, total = B._rows(self.OK_XML)
        self.assertEqual(total, 2)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["연면적"], "5230.5")
        self.assertEqual(rows[0]["사용승인일"], "19980312")

    def test_키가_안_먹으면_빈_결과다(self):
        """서비스 키 미등록은 HTTP 200 에 오류 코드로 온다. 이걸 자료로 세면
        연면적 0 ㎡ 가 서식에 찍힌다."""
        from firebird import buildings as B
        bad = ("<response><header><resultCode>30</resultCode>"
               "<resultMsg>SERVICE KEY IS NOT REGISTERED ERROR</resultMsg>"
               "</header></response>")
        self.assertEqual(B._rows(bad), ([], 0))

    def test_깨진_응답에도_죽지_않는다(self):
        from firebird import buildings as B
        for junk in ("", "not xml", "<response>"):
            self.assertEqual(B._rows(junk), ([], 0))

    def test_격자에_값이_없으면_빈_사전(self):
        import pandas as pd
        from firebird import buildings as B
        df = pd.DataFrame([{"grid_id": "1_1", "연면적": 100.0, "건축면적": 20.0,
                            "사용승인연도": 2001.0, "건물수": 3}])
        self.assertEqual(B.stats_for(df, "9_9"), {})
        self.assertEqual(B.stats_for(pd.DataFrame(), "1_1"), {})
        got = B.stats_for(df, "1_1")
        self.assertEqual(got["연면적"], "100")
        self.assertEqual(got["건축연도"], "2001")

    def test_건축물대장이_없으면_칸이_비고_경로가_남는다(self):
        import pandas as pd
        from firebird import forms as FM
        row = pd.Series({"grid_id": "1_1", "lon": 129.3, "lat": 35.5,
                         "target_total": 10, "biz_total": 3, "fires": 1,
                         "emd": "달동", "sgg": "남구"})
        led = FM.zone_ledger(row, city_label="울산광역시", year=2021)
        f = led["fields"]["연면적"]
        self.assertEqual(f.value, "")
        self.assertIn("건축물대장", f.blank_reason)
        self.assertIn("건축물대장", f.route)

    def test_건축물대장이_있으면_칸이_채워지고_출처가_남는다(self):
        import pandas as pd
        from firebird import forms as FM
        row = pd.Series({"grid_id": "1_1", "lon": 129.3, "lat": 35.5,
                         "target_total": 10, "biz_total": 3, "fires": 1,
                         "emd": "달동", "sgg": "남구"})
        led = FM.zone_ledger(row, city_label="울산광역시", year=2021,
                             building={"연면적": "6,350", "건축연도": "2006", "_n": 2})
        f = led["fields"]["연면적"]
        self.assertEqual(f.value, "6,350")
        self.assertEqual(f.blank_reason, "")
        self.assertIn("건축물대장", f.source)
        # 안 온 칸은 그대로 비어 있어야 한다
        self.assertEqual(led["fields"]["건축면적"].value, "")

    def test_건물동수는_대상물_수가_아니다(self):
        """서식의 '건물동수' 는 건축물 동수다. 특정소방대상물은 대상물 단위라
        한 건물에 여러 건이 등록될 수 있어 동수보다 크게 나온다
        (실측: 한 격자에서 대상물 413건 vs 건축물 206동). 둘을 같은 칸에
        넣으면 서식이 묻는 값과 다른 값이 들어간다."""
        import pandas as pd
        from firebird import forms as FM
        row = pd.Series({"grid_id": "1_1", "lon": 129.3, "lat": 35.5,
                         "target_total": 413, "biz_total": 159, "fires": 1,
                         "emd": "달동", "sgg": "남구"})
        # 건축물대장이 없으면 그 칸은 비운다 — 대상물 수를 대신 넣지 않는다
        led = FM.zone_ledger(row, city_label="울산광역시", year=2021)
        self.assertEqual(led["fields"]["건물동수"].value, "")
        self.assertIn("건축물대장", led["fields"]["건물동수"].route)
        # 대상물 수는 지구특징에 남는다
        self.assertIn("413", led["fields"]["지구특징"].value)

        led2 = FM.zone_ledger(row, city_label="울산광역시", year=2021,
                              building={"연면적": "380,976", "_n": 206})
        self.assertEqual(led2["fields"]["건물동수"].value, "206")
        self.assertIn("건축물대장", led2["fields"]["건물동수"].source)


class TestPopulationApi(unittest.TestCase):
    """SGIS 인구 연계.

    가장 중요한 것: **읍면동 값을 격자 값인 척 쓰지 않는다.** 서식의 상주인구
    칸은 비워 두고, 읍면동 단위임을 이름에 박아 다른 곳에서만 쓴다.
    """

    def test_SGIS_코드는_법정동코드가_아니다(self):
        """법정동코드 31140 은 울산 남구지만 SGIS 31140 은 오산시다.
        코드를 그대로 넘기면 다른 도시 인구가 들어온다."""
        from firebird import population as P
        self.assertIn("법정동코드", P.__doc__)
        self.assertTrue(hasattr(P, "find_code"))

    def test_상주인구_칸은_비운다(self):
        import pandas as pd
        from firebird import forms as FM
        row = pd.Series({"grid_id": "1_1", "lon": 129.3, "lat": 35.5,
                         "target_total": 10, "biz_total": 3, "fires": 1,
                         "emd": "달동", "sgg": "남구"})
        led = FM.zone_ledger(row, city_label="울산광역시", year=2021)
        f = led["fields"]["상주인구"]
        self.assertEqual(f.value, "")
        self.assertIn("격자", f.blank_reason)

    def test_인구를_붙이면_읍면동_단위임이_이름에_남는다(self):
        import pandas as pd
        from firebird import population as P
        grid = pd.DataFrame([{"grid_id": "1_1", "sgg": "남구", "emd": "달동"}])
        pop = pd.DataFrame([{"sgg": "남구", "emd": "달동", "상주인구": 26067,
                             "인구밀도": 20149.0}])
        out = P.attach(grid, pop)
        self.assertIn("읍면동 인구", out.columns)
        self.assertNotIn("상주인구", out.columns)
        self.assertEqual(int(out.iloc[0]["읍면동 인구"]), 26067)

    def test_키가_없으면_빈_표를_돌려준다(self):
        import pandas as pd
        from firebird import population as P
        self.assertTrue(P.attach(pd.DataFrame(), pd.DataFrame()).empty)
        grid = pd.DataFrame([{"grid_id": "1_1", "sgg": "남구", "emd": "달동"}])
        pd.testing.assert_frame_equal(P.attach(grid, pd.DataFrame()), grid)


class TestBuildingAgeFeatures(unittest.TestCase):
    """연도별 노후도 피처.

    이 피처의 존재 이유는 '연도에 따라 변한다'는 것 하나다. 스냅샷 피처처럼
    모든 연도에 같은 값이 들어가면 만든 의미가 없다.
    """

    def _b(self):
        import pandas as pd
        return pd.DataFrame([
            {"지번주소": "울산광역시 남구 달동 100", "연면적": 1000.0,
             "사용승인연도": 1990.0},
            {"지번주소": "울산광역시 남구 달동 101", "연면적": 2000.0,
             "사용승인연도": 2018.0},
            {"지번주소": "울산광역시 남구 달동 102", "연면적": 500.0,
             "사용승인연도": 2019.0},
            # 연도가 없는 건물은 어느 해에도 세지 않는다
            {"지번주소": "울산광역시 남구 달동 103", "연면적": 900.0,
             "사용승인연도": None},
            # 범위 밖 값(원본에 978년, 2026년 같은 것이 섞여 있다)
            {"지번주소": "울산광역시 남구 달동 104", "연면적": 100.0,
             "사용승인연도": 978.0},
        ])

    def test_해마다_값이_달라진다(self):
        from firebird import buildings as B
        f = B.emd_year_features(None, [2015, 2017, 2018, 2020],
                                buildings=self._b())
        by = f.set_index("year")
        self.assertEqual(int(by.loc[2015, "bld_n"]), 1)      # 1990년 건물만
        self.assertEqual(int(by.loc[2018, "bld_n"]), 2)
        self.assertEqual(int(by.loc[2020, "bld_n"]), 3)
        # 건물 구성이 그대로면 평균 연령은 해마다 1씩 오른다
        self.assertAlmostEqual(float(by.loc[2017, "bld_age_mean"])
                               - float(by.loc[2015, "bld_age_mean"]), 2.0)
        # 신축이 들어오면 평균 연령은 오히려 내려간다 — 그게 맞는 동작이다
        self.assertLess(by.loc[2018, "bld_age_mean"], by.loc[2017, "bld_age_mean"])

    def test_미래_준공은_세지_않는다(self):
        """t년 피처에 t년 이후 준공 건물이 들어가면 그게 곧 누수다."""
        from firebird import buildings as B
        f = B.emd_year_features(None, [2015], buildings=self._b())
        self.assertEqual(float(f.iloc[0]["bld_area"]), 1000.0)

    def test_범위_밖_연도와_결측은_버린다(self):
        from firebird import buildings as B
        f = B.emd_year_features(None, [2021], buildings=self._b())
        self.assertEqual(int(f.iloc[0]["bld_n"]), 3)

    def test_노후비율은_0과_1_사이(self):
        from firebird import buildings as B
        f = B.emd_year_features(None, [2015, 2021], buildings=self._b())
        for c in ("bld_old_share", "bld_new_share"):
            self.assertTrue(((f[c] >= 0) & (f[c] <= 1)).all())

    def test_붙일_자료가_없으면_패널을_그대로_둔다(self):
        import pandas as pd
        from firebird import buildings as B
        panel = pd.DataFrame([{"grid_id": "1_1", "sgg": "남구", "emd": "달동",
                               "year": 2021, "fires": 1}])
        pd.testing.assert_frame_equal(B.attach_features(panel, pd.DataFrame()),
                                      panel)

    def test_측정_결과가_기록되어_있다(self):
        """넣을지 말지는 측정으로 정한다. 그 측정이 파일로 남아 있어야
        나중에 '왜 안 넣었나'에 답할 수 있다."""
        import json
        from pathlib import Path
        p = (Path(__file__).resolve().parents[1] / "outputs"
             / "building_feature_eval_ulsan.json")
        if not p.exists():
            self.skipTest("scripts/16_eval_buildings.py 를 아직 돌리지 않았다")
        got = json.loads(p.read_text(encoding="utf-8"))
        self.assertIn("adopted", got)
        self.assertIn("delta_ci", got)
        self.assertIn("history_only", got)


class TestStreamlitWidgetState(unittest.TestCase):
    """위젯 상태를 세션에 넣을 때 키를 맞췄는지 본다.

    '자주 찾는 질문' 버튼이 눌려도 아무 일이 없던 적이 있다. 버튼이
    st.session_state["qa_question"] 만 채웠는데, 정작 질문을 읽는 위젯의
    키는 "qa_input" 이었다. text_area 의 value= 는 첫 렌더의 기본값일 뿐
    이미 만들어진 위젯에는 반영되지 않는다. 시연 중에 드러나면 늦다.
    """

    def setUp(self):
        from pathlib import Path
        self.src = (Path(__file__).resolve().parents[1] / "app"
                    / "streamlit_app.py").read_text(encoding="utf-8")

    def test_자주찾는질문_버튼이_위젯_키를_채운다(self):
        self.assertIn('st.session_state["qa_input"] = q', self.src)

    def test_질문_위젯_키가_그대로다(self):
        self.assertIn('key="qa_input"', self.src)


class TestDeckHasNoTrailingNotes(unittest.TestCase):
    """장표 아래에 회색 한 줄을 매달아 그림을 다시 설명하지 않는지 본다.

    '검은 점 = 출동 관서', '두 지도는 같은 범위입니다' 처럼 그림을 보면 아는
    것을 밑에 또 적어 두는 버릇이 있었다. 발표에서 말하면 되는 것이고,
    장표에 있으면 AI 가 덧붙인 티가 난다. 한 번 걷어낸 뒤 다시 늘지 않게 한다.
    """

    def test_그림_아래_회색_설명줄이_없다(self):
        from pathlib import Path
        try:
            from pptx import Presentation
        except ImportError:                               # noqa: BLE001
            self.skipTest("python-pptx 없음")
        root = Path(__file__).resolve().parents[1]
        found = list((root / "outputs").glob("불씨예보_발표자료_*.pptx"))
        if not found:
            self.skipTest("발표자료를 아직 만들지 않았다")
        muted = (0x6B, 0x74, 0x84)
        leftovers = []
        for i, slide in enumerate(Presentation(found[0]).slides, 1):
            for sh in slide.shapes:
                if not sh.has_text_frame or not sh.text_frame.text.strip():
                    continue
                wide = bool(sh.width and sh.width.inches > 10)
                low = bool(sh.top and sh.top.inches > 5.9)
                runs = sh.text_frame.paragraphs[0].runs
                if not (wide and low and runs):
                    continue
                col = runs[0].font.color
                try:
                    rgb = tuple(col.rgb) if col and col.type is not None else None
                except Exception:                         # noqa: BLE001
                    rgb = None
                if rgb == muted:
                    leftovers.append((i, sh.text_frame.text.strip()[:60]))
        self.assertEqual(leftovers, [], f"장표 아래 회색 설명줄: {leftovers}")


class TestDeckTitleStyle(unittest.TestCase):
    """장표 제목이 말하듯 한 문장이 아니라 명사구인지 본다.

    '저희가 내놓는 것은 이 문서입니다' 같은 제목이 섞여 있었다. 공모전
    발표자료는 개조식 명사구가 관행이고, 문장형과 명사구가 섞이면 그것부터
    눈에 띈다. 한 번 맞춘 문체가 다음 판에서 흐트러지지 않게 한다.
    """

    ENDINGS = ("니다", "습니다", "한다", "된다", "입니다", "짜입니다")

    def test_제목이_문장형이_아니다(self):
        from pathlib import Path
        try:
            from pptx import Presentation
        except ImportError:                               # noqa: BLE001
            self.skipTest("python-pptx 없음")
        root = Path(__file__).resolve().parents[1]
        found = list((root / "outputs").glob("불씨예보_발표자료_*.pptx"))
        if not found:
            self.skipTest("발표자료를 아직 만들지 않았다")
        bad = []
        for i, slide in enumerate(Presentation(found[0]).slides, 1):
            for sh in slide.shapes:
                if not sh.has_text_frame:
                    continue
                big = any(r.font.size and r.font.size.pt >= 26
                          for para in sh.text_frame.paragraphs for r in para.runs)
                if not big:
                    continue
                text = sh.text_frame.text.strip().replace("\n", " ")
                if text.endswith(self.ENDINGS):
                    bad.append((i, text[:50]))
        self.assertEqual(bad, [], f"문장형 제목: {bad}")


class TestPlanPreviewHtml(unittest.TestCase):
    """계획서 미리보기가 문서를 있는 그대로 그리는가.

    결재란의 서명 칸이 화면에서 사라져 있었다. 빈 행(| | | |)의 셀이 전부
    비어 있어 all([]) 이 True 가 되고, 구분선으로 오인돼 통째로 지워졌다.
    관서로 나가는 문서라 빈칸도 문서의 일부다.
    """

    @staticmethod
    def _fns():
        import ast
        from pathlib import Path
        src = Path(__file__).resolve().parents[1] / "app" / "streamlit_app.py"
        tree = ast.parse(src.read_text(encoding="utf-8"))
        want = {"_fmt_inline", "_md_to_html", "_img_src"}
        mod = ast.Module(body=[n for n in tree.body
                               if isinstance(n, ast.FunctionDef) and n.name in want],
                         type_ignores=[])
        ns: dict = {}
        exec(compile(mod, "<app>", "exec"), ns)  # noqa: S102
        return ns

    def test_empty_rows_survive(self):
        html = self._fns()["_md_to_html"](
            "| 기안 | 검토 | 결재 |\n|---|---|---|\n| | | |\n| | | |")
        self.assertEqual(html.count("<tr>"), 3, f"빈 행이 사라졌다:\n{html}")

    def test_separator_row_still_dropped(self):
        html = self._fns()["_md_to_html"]("| 가 | 나 |\n|---|:--:|\n| 1 | 2 |")
        self.assertNotIn("---", html)
        self.assertEqual(html.count("<tr>"), 2)

    def test_local_image_is_inlined(self):
        import tempfile, os
        fns = self._fns()
        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as fh:
            fh.write(b"\x89PNG\r\n\x1a\n" + b"\0" * 64)
            path = fh.name
        try:
            html = fns["_md_to_html"](f"![동선도]({path})")
            self.assertIn("data:image/png;base64,", html,
                          "화면 미리보기가 파일 경로를 그대로 넣고 있다")
        finally:
            os.unlink(path)
